"""
Telegram-бот клуба COLIZEUM.

Что делает:
- /start — присылает кнопку, открывающую мини-приложение (личный кабинет).
- Админ-команды (доступны только tg_id из ADMIN_IDS) — временный способ
  вести данные о пополнениях и акциях, пока нет автоматической синхронизации
  с CRM/POS клуба. Как только появится выгрузка или API, эти команды можно
  будет заменить на автоматический импорт, ничего в остальном коде менять
  не придётся.

Бот работает поверх той же базы данных, что и веб-API мини-приложения
(см. main.py) — это один процесс, поэтому отдельный сервис под бота
разворачивать не нужно.
"""
import asyncio
import logging
import os
import secrets
from datetime import datetime, timedelta
from io import BytesIO
from urllib.parse import urlsplit, urlunsplit, parse_qsl, urlencode

from aiogram import Bot, Dispatcher, F, Router
from aiogram.exceptions import TelegramForbiddenError, TelegramRetryAfter, TelegramBadRequest
from aiogram.filters import Command, CommandObject
from aiogram.types import (
    Message, CallbackQuery, WebAppInfo, InlineKeyboardMarkup, InlineKeyboardButton,
    ReplyKeyboardMarkup, KeyboardButton, ReplyKeyboardRemove, BotCommand,
)
from sqlalchemy import or_, func

from .database import SessionLocal
from .models import (
    Client, Promotion, PhoneStat, ReferralReward, GamePlay, get_tier, normalize_phone,
    lookup_topup, load_topup_map, topup_key, resync_client_topups,
    TIER_LABELS, SILVER_QUALIFY_FROM,
)
from .topups import parse_topups
from . import referrals

logger = logging.getLogger("colizeum")

# Меню команд, которое гости видят в боте (кнопка «Меню»). Админ-команды
# сюда не попадают — гостям они не нужны и всё равно не сработают.
GUEST_COMMANDS = [
    BotCommand(command="start", description="Открыть личный кабинет"),
    BotCommand(command="bonus", description="🎁 Пригласить друга и получить бонус"),
    BotCommand(command="phone", description="Подтвердить номер телефона"),
    BotCommand(command="stop", description="Отписаться от рассылки"),
]

WEBAPP_URL = os.getenv("WEBAPP_URL", "")
ADMIN_IDS = {
    int(x) for x in os.getenv("ADMIN_IDS", "").replace(" ", "").split(",") if x
}

router = Router()


def is_admin(tg_id: int) -> bool:
    return tg_id in ADMIN_IDS


def humanize_ago(dt) -> str:
    """Переводит момент времени в «X назад» на русском — для списка гостей."""
    if not dt:
        return "не заходил(а)"
    delta = datetime.utcnow() - dt
    minutes = int(delta.total_seconds() // 60)
    if minutes < 1:
        return "только что"
    if minutes < 60:
        return f"{minutes} мин назад"
    hours = minutes // 60
    if hours < 24:
        return f"{hours} ч назад"
    days = hours // 24
    return f"{days} дн назад"


def webapp_url(ref: str | None = None) -> str:
    """Адрес мини-приложения; при переходе по ссылке друга — с меткой ref."""
    if not ref:
        return WEBAPP_URL
    parts = urlsplit(WEBAPP_URL)
    query = dict(parse_qsl(parts.query))
    query["ref"] = ref
    return urlunsplit((parts.scheme, parts.netloc, parts.path, urlencode(query), parts.fragment))


def cabinet_button(ref: str | None = None) -> InlineKeyboardButton:
    return InlineKeyboardButton(text="Открыть личный кабинет", web_app=WebAppInfo(url=webapp_url(ref)))


@router.message(Command("start"))
async def cmd_start(message: Message, command: CommandObject):
    if not WEBAPP_URL:
        await message.answer(
            "Мини-приложение ещё не подключено (не задан WEBAPP_URL). "
            "Как только опубликуешь webapp, добавь адрес в настройки — и эта кнопка заработает."
        )
        return

    arg = (command.args or "").strip()
    ref = arg if referrals.parse_ref(arg) is not None else None
    if arg == "bonus":
        await send_bonus_info(message.bot, message.from_user.id)
        return

    keyboard = InlineKeyboardMarkup(inline_keyboard=[
        [cabinet_button(ref)],
        [InlineKeyboardButton(text="🎁 Получить бонус", callback_data="bonus")],
    ])
    intro = ("Тебя пригласил друг — добро пожаловать в COLIZEUM!\n\n" if ref
             else "Добро пожаловать в COLIZEUM.\n\n")
    await message.answer(
        intro +
        "Здесь — твой уровень и кэшбэк, акции клуба и мини-игры с призами. "
        "Открывай, когда удобно.",
        reply_markup=keyboard,
    )


# ---------- «Получить бонус»: пригласи друга ----------

async def send_bonus_info(bot: Bot, tg_id: int):
    """Личная ссылка для друзей. Пишем гостю в личный чат с ботом."""
    db = SessionLocal()
    try:
        client = db.query(Client).filter(Client.tg_id == tg_id).first()
        if client is None or client.consent_at is None:
            await bot.send_message(
                tg_id,
                "Сначала открой личный кабинет и прими правила — после этого здесь появится "
                "твоя личная ссылка для друзей.",
                reply_markup=InlineKeyboardMarkup(inline_keyboard=[[cabinet_button()]]) if WEBAPP_URL else None,
            )
            return
        invited = referrals.invited_count(db, client)
        client_id = client.id
    finally:
        db.close()

    username = referrals.BOT_USERNAME or (await bot.me()).username
    link = referrals.referral_link(client_id, username)
    bonus = referrals.REFERRAL_BONUS
    await bot.send_message(
        tg_id,
        f"🎁 Приведи друга в бота — получи {bonus} бонусов\n\n"
        "1. Нажми «Отправить другу» и выбери друга в Telegram.\n"
        "2. Друг открывает бота, принимает правила и подтверждает номер.\n"
        f"3. Тебе сразу приходят +{bonus} бонусов на баланс.\n\n"
        "Засчитываются только новые пользователи бота, за каждый номер телефона — один раз.\n\n"
        f"Твоя ссылка: {link}\n"
        f"Друзей уже привёл: {invited}",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[[
            InlineKeyboardButton(text="📤 Отправить другу", url=referrals.share_url(link))
        ]]),
        disable_web_page_preview=True,
    )


@router.message(Command("bonus"))
async def cmd_bonus(message: Message):
    await send_bonus_info(message.bot, message.from_user.id)


@router.callback_query(F.data == "bonus")
async def cb_bonus(callback: CallbackQuery):
    await callback.answer()
    await send_bonus_info(callback.bot, callback.from_user.id)


# ---------- отписка от рассылки ----------

def _set_opt_out(tg_id: int, value: bool) -> bool:
    db = SessionLocal()
    try:
        client = db.query(Client).filter(Client.tg_id == tg_id).first()
        if client is None:
            return False
        client.broadcast_opt_out = value
        db.commit()
        return True
    finally:
        db.close()


@router.message(Command("stop"))
async def cmd_stop(message: Message):
    _set_opt_out(message.from_user.id, True)
    await message.answer("Готово — рассылки от клуба больше не придут. "
                         "Личный кабинет работает как раньше. Вернуть рассылку: /subscribe")


@router.message(Command("subscribe"))
async def cmd_subscribe(message: Message):
    if _set_opt_out(message.from_user.id, False):
        await message.answer("Готово — снова будешь получать новости и акции клуба.")
    else:
        await message.answer("Сначала открой личный кабинет: /start")


def save_phone_stats(db, stats: dict, replace_all: bool = False) -> tuple[int, int]:
    """Сохраняет средние пополнения по телефонам и сразу проставляет их
    гостям, которые уже привязали номер.

    replace_all=True — для полной выгрузки из CRM: статистика заменяется
    целиком. Телефоны, которых нет в новой выгрузке (гость перестал ходить),
    обнуляются, чтобы старый статус не висел вечно.
    Возвращает (скольким гостям в приложении обновлён уровень, скольким обнулён)."""
    now = datetime.utcnow()
    reset = 0
    if replace_all:
        db.query(PhoneStat).delete(synchronize_session=False)
        stale = db.query(Client).filter(Client.phone_normalized.isnot(None),
                                        Client.monthly_topup > 0).all()
        for client in stale:
            if client.phone_normalized not in stats:
                client.monthly_topup = 0.0
                reset += 1
        db.flush()
    phones = list(stats)
    chunks = [phones[i:i + 500] for i in range(0, len(phones), 500)]  # лимит параметров в SQL
    existing = {}
    for chunk in chunks:
        for s in db.query(PhoneStat).filter(PhoneStat.phone_normalized.in_(chunk)).all():
            existing[s.phone_normalized] = s
    for phone, st in stats.items():
        row = existing.get(phone)
        if row is None:
            db.add(PhoneStat(phone_normalized=phone, avg_monthly=st["avg"],
                             total=st.get("total"), updated_at=now))
        else:
            row.avg_monthly = st["avg"]
            row.total = st.get("total")
            row.updated_at = now

    db.commit()
    # Пересчитываем уровни всех гостей приложения по новой статистике
    # (учитывает и тех, кого связали с другим номером через /link_crm).
    resync_client_topups(db)
    updated = sum(1 for c in db.query(Client).filter(Client.phone_verified == True).all()  # noqa: E712
                  if topup_key(c) in stats)
    return updated, reset


PHONE_KEYBOARD = ReplyKeyboardMarkup(
    keyboard=[[KeyboardButton(text="📱 Подтвердить номер", request_contact=True)]],
    resize_keyboard=True, one_time_keyboard=True,
)


@router.message(Command("phone"))
async def cmd_phone(message: Message):
    """Запасной способ подтвердить номер, если в мини-приложении кнопка не сработала."""
    await message.answer(
        "Нажми кнопку ниже — Telegram передаст номер, привязанный к твоему аккаунту.",
        reply_markup=PHONE_KEYBOARD,
    )


@router.message(F.contact)
async def handle_contact(message: Message):
    """Гость поделился номером (из мини-приложения или кнопкой /phone).
    Номер приходит от самого Telegram — подделать его нельзя. Принимаем
    только СВОЙ контакт: переслать чужой не выйдет."""
    contact = message.contact
    if contact.user_id != message.from_user.id:
        await message.answer("Подтвердить можно только свой номер — нажми кнопку «📱 Подтвердить номер».",
                             reply_markup=PHONE_KEYBOARD)
        return

    digits = "".join(ch for ch in contact.phone_number if ch.isdigit())
    norm = normalize_phone(digits)
    db = SessionLocal()
    try:
        client = db.query(Client).filter(Client.tg_id == message.from_user.id).first()
        if client is None or client.consent_at is None:
            await message.answer(
                "Сначала открой личный кабинет (/start) и дай согласие на обработку данных, "
                "потом подтверди номер.", reply_markup=ReplyKeyboardRemove())
            return

        # Один номер — один аккаунт: если номер был у другого аккаунта, отвязываем его там.
        others = db.query(Client).filter(Client.phone_normalized == norm, Client.id != client.id).all()
        for other in others:
            other.phone = None
            other.phone_normalized = None
            other.phone_verified = False
            other.monthly_topup = 0.0

        client.phone = "+" + digits
        client.phone_normalized = norm
        client.phone_verified = True
        topup, crm_found = lookup_topup(db, client)
        client.monthly_topup = topup
        db.commit()
        tier = TIER_LABELS[get_tier(topup)]
        # Номер уже был в боте у другого аккаунта — это не новый гость, бонус за друга не платим.
        reward = referrals.reward_referral(db, client) if not others else None
    finally:
        db.close()

    tail = "" if crm_found else (
        "\n\nПо этому номеру не нашли пополнений в клубе. Если в клубе ты зарегистрирован "
        "на другой номер — скажи администратору, он свяжет номера.")
    await message.answer(
        f"✅ Номер подтверждён. Твой уровень: {tier}.\n"
        "Возвращайся в личный кабинет — игры и списание бонусов теперь доступны." + tail,
        reply_markup=ReplyKeyboardRemove(),
    )

    if reward and reward[1] > 0:
        referrer_tg_id, amount = reward
        friend = message.from_user.first_name or "Друг"
        try:
            await message.bot.send_message(
                referrer_tg_id,
                f"🎉 {friend} зарегистрировался по твоей ссылке — +{amount} бонусов на баланс!\n"
                "Пригласить ещё: /bonus")
        except Exception as e:  # noqa: BLE001 — пригласивший мог заблокировать бота
            logger.info("Не удалось уведомить пригласившего %s: %s", referrer_tg_id, e)


@router.message(Command("set_topup"))
async def cmd_set_topup(message: Message, command: CommandObject):
    """Формат: /set_topup +79991234567 12000
    Вручную задаёт средние пополнения в месяц по номеру. Работает и для
    гостей, которые ещё не заходили в приложение — уровень появится у них
    при привязке номера."""
    if not is_admin(message.from_user.id):
        return

    parts = (command.args or "").split()
    if len(parts) != 2:
        await message.answer("Формат: /set_topup +79991234567 12000")
        return

    phone, amount_raw = parts
    try:
        amount = float(amount_raw.replace(",", "."))
    except ValueError:
        await message.answer("Сумма должна быть числом, например 12000")
        return

    norm = normalize_phone(phone)
    if len(norm) < 10:
        await message.answer("Не похоже на номер телефона — нужно 10 цифр после +7")
        return

    db = SessionLocal()
    try:
        updated, _ = save_phone_stats(db, {norm: {"avg": amount, "total": None}})
    finally:
        db.close()

    tier = TIER_LABELS[get_tier(amount)]
    tail = "Гость уже в приложении — уровень обновлён." if updated else \
        "Гость ещё не привязал номер — уровень появится сразу при привязке."
    await message.answer(f"Готово: {phone} — {amount:,.0f} ₽ в месяц, уровень {tier}.\n{tail}".replace(",", " "))


@router.message(Command("add_promo"))
async def cmd_add_promo(message: Message, command: CommandObject):
    """Формат: /add_promo Название | Описание | ссылка (необязательно)
    Акции видны всем гостям клуба независимо от уровня."""
    if not is_admin(message.from_user.id):
        return

    if not command.args:
        await message.answer("Формат: /add_promo Название | Описание | ссылка (необязательно)")
        return

    parts = [p.strip() for p in command.args.split("|")]
    if len(parts) < 2:
        await message.answer("Формат: /add_promo Название | Описание | ссылка (необязательно)")
        return

    title, description = parts[0], parts[1]
    link = parts[2] if len(parts) > 2 and parts[2] else None

    db = SessionLocal()
    try:
        promo = Promotion(title=title, description=description, link=link)
        db.add(promo)
        db.commit()
        await message.answer(f"Акция добавлена: «{title}»" + (f"\nСсылка: {link}" if link else ""))
    finally:
        db.close()


@router.message(Command("promotions"))
async def cmd_list_promotions(message: Message):
    if not is_admin(message.from_user.id):
        return

    db = SessionLocal()
    try:
        promos = db.query(Promotion).filter(Promotion.active == True).all()  # noqa: E712
        if not promos:
            await message.answer("Активных акций пока нет. Добавь: /add_promo Название | Описание")
            return
        lines = [f"#{p.id} {p.title} — {p.description}" for p in promos]
        await message.answer("\n".join(lines))
    finally:
        db.close()


@router.message(Command("stats"))
async def cmd_stats(message: Message):
    """Сводка: сколько всего гостей в базе, разбивка по уровням,
    сколько заходило за последние сутки и за последнюю неделю."""
    if not is_admin(message.from_user.id):
        return

    db = SessionLocal()
    try:
        clients = db.query(Client).all()
        total = len(clients)
        if total == 0:
            await message.answer("Гостей пока нет — база пустая.")
            return

        topups = load_topup_map(db)
        by_tier = {name: 0 for name in TIER_LABELS}
        active_24h = 0
        active_7d = 0
        verified = 0
        found_in_crm = 0
        opted_out = 0
        now = datetime.utcnow()
        for c in clients:
            key = topup_key(c)
            by_tier[get_tier(topups.get(key, 0) if key else 0)] += 1
            verified += 1 if c.phone_verified else 0
            found_in_crm += 1 if key and key in topups else 0
            opted_out += 1 if c.broadcast_opt_out else 0
            if c.last_seen_at:
                delta_hours = (now - c.last_seen_at).total_seconds() / 3600
                if delta_hours <= 24:
                    active_24h += 1
                if delta_hours <= 24 * 7:
                    active_7d += 1

        lines = [
            f"Всего гостей в базе: {total}",
            f"Заходили за 24 часа: {active_24h}",
            f"Заходили за 7 дней: {active_7d}",
            "",
            "По уровням:",
            f"  Premium — {by_tier['premium']}",
            f"  Gold — {by_tier['gold']}",
            f"  Silver — {by_tier['silver']}",
            f"  Без статуса — {by_tier['none']}",
            "",
            f"Подтвердили номер: {verified} из {total}",
            f"  из них найдены в CRM: {found_in_crm}, не найдены: {verified - found_in_crm}",
            f"Пришли по приглашению друзей: {db.query(ReferralReward).count()}",
            f"Отписались от рассылки: {opted_out}",
        ]
        # Диагностика: загружена ли статистика из CRM и когда.
        stats_count = db.query(PhoneStat).count()
        last = db.query(PhoneStat.updated_at).order_by(PhoneStat.updated_at.desc()).first()
        if stats_count:
            lines.append(f"Статистика CRM: {stats_count} телефонов, загружена "
                         f"{(last[0] + timedelta(hours=3)):%d.%m.%Y %H:%M} (МСК)")
        else:
            lines.append("⚠️ Статистика CRM НЕ загружена — уровни ни у кого не появятся. "
                         "Отправь лог операций с подписью /import_topups")
        await message.answer("\n".join(lines))
    finally:
        db.close()


@router.message(Command("check"))
async def cmd_check(message: Message, command: CommandObject):
    """Формат: /check +79991234567
    Показывает всё по номеру: что о нём известно из CRM и какие аккаунты
    приложения к нему привязаны. Для разбора «почему у гостя нет статуса»."""
    if not is_admin(message.from_user.id):
        return
    norm = normalize_phone(command.args or "")
    if len(norm) < 10:
        await message.answer("Формат: /check +79991234567")
        return

    db = SessionLocal()
    try:
        stat = db.get(PhoneStat, norm)
        linked = db.query(Client).filter(or_(Client.phone_normalized == norm,
                                             Client.crm_phone_normalized == norm)).all()
        lines = [f"Номер +7{norm}", ""]
        if stat:
            tier = TIER_LABELS[get_tier(stat.avg_monthly or 0)]
            lines.append(f"В статистике CRM: да — {round(stat.avg_monthly or 0):,} ₽/мес в среднем → {tier}"
                         .replace(",", " "))
        else:
            lines.append("В статистике CRM: НЕТ (гость не пополнял счёт за период выгрузки "
                         "или статистика не загружена — см. /stats)")
        lines.append("")
        if not linked:
            lines.append("В приложении этот номер никто не подтвердил.")
        for c in linked:
            topup, _ = lookup_topup(db, c)
            crm = f", уровень считается по номеру CRM +7{c.crm_phone_normalized}" if c.crm_phone_normalized else ""
            lines.append(
                f"Аккаунт: {c.tg_name or 'без имени'} (id {c.tg_id}), Telegram-номер {c.phone or '—'} — "
                f"{'номер подтверждён' if c.phone_verified else 'номер НЕ подтверждён'}, "
                f"уровень {TIER_LABELS[get_tier(topup)]}, баланс {c.balance or 0}{crm}"
            )
        await message.answer("\n".join(lines))
    finally:
        db.close()


@router.message(Command("clients"))
async def cmd_clients(message: Message):
    """Список последних гостей по времени визита — кто и когда заходил.
    Показывает 20 самых недавних; для точной сверки сумм используй выгрузку из CRM."""
    if not is_admin(message.from_user.id):
        return

    db = SessionLocal()
    try:
        clients = (
            db.query(Client)
            .order_by(Client.last_seen_at.desc().nullslast())
            .limit(20)
            .all()
        )
        if not clients:
            await message.answer("Гостей пока нет — база пустая.")
            return

        topups = load_topup_map(db)
        lines = ["Последние 20 гостей (по времени визита):", ""]
        for c in clients:
            name = c.tg_name or "без имени"
            phone = c.phone or "телефон не привязан"
            key = topup_key(c)
            tier = TIER_LABELS[get_tier(topups.get(key, 0) if key else 0)]
            lines.append(f"{name} — {phone} — {tier} — {humanize_ago(c.last_seen_at)}")
        await message.answer("\n".join(lines))
    finally:
        db.close()


@router.message(Command("spend"))
async def cmd_spend(message: Message, command: CommandObject):
    """Формат: /spend +79991234567 150
    Списывает бонусы с баланса гостя, когда он расплачивается ими на кассе."""
    if not is_admin(message.from_user.id):
        return

    if not command.args:
        await message.answer("Формат: /spend +79991234567 150")
        return

    parts = command.args.split()
    if len(parts) != 2:
        await message.answer("Формат: /spend +79991234567 150")
        return

    phone, amount_raw = parts
    try:
        amount = int(float(amount_raw))
    except ValueError:
        await message.answer("Сумма должна быть числом, например 150")
        return

    if amount <= 0:
        await message.answer("Сумма списания должна быть больше нуля")
        return

    db = SessionLocal()
    try:
        client = (db.query(Client).filter(Client.phone_normalized == normalize_phone(phone))
                  .order_by(Client.phone_verified.desc()).first())
        if not client:
            await message.answer(f"Клиент с номером {phone} не найден — он ещё не заходил в приложение.")
            return
        current = client.balance or 0
        if current < amount:
            await message.answer(
                f"Недостаточно бонусов: на балансе {current}, а списать нужно {amount}."
            )
            return
        client.balance = current - amount
        db.commit()
        await message.answer(f"Списано {amount} бонусов. Остаток на балансе: {client.balance}.")
    finally:
        db.close()


@router.message(Command("add_bonus"))
async def cmd_add_bonus(message: Message, command: CommandObject):
    """Формат: /add_bonus +79991234567 150
    Начисляет бонусы на баланс в приложении — например, вернуть бонусы
    по отклонённой заявке на списание или начислить вручную."""
    if not is_admin(message.from_user.id):
        return

    parts = (command.args or "").split()
    if len(parts) != 2:
        await message.answer("Формат: /add_bonus +79991234567 150")
        return

    phone, amount_raw = parts
    try:
        amount = int(float(amount_raw))
    except ValueError:
        await message.answer("Сумма должна быть числом, например 150")
        return
    if amount <= 0:
        await message.answer("Сумма должна быть больше нуля")
        return

    db = SessionLocal()
    try:
        client = (db.query(Client).filter(Client.phone_normalized == normalize_phone(phone))
                  .order_by(Client.phone_verified.desc()).first())
        if not client:
            await message.answer(
                f"Клиент с номером {phone} не найден — номер должен быть привязан в приложении."
            )
            return
        client.balance = (client.balance or 0) + amount
        db.commit()
        await message.answer(f"Начислено {amount} бонусов. Баланс в приложении: {client.balance}.")
    finally:
        db.close()


@router.message(F.document)
async def handle_document(message: Message):
    """Загрузка пополнений. Прикрепи в Telegram файл Excel (.xlsx) с подписью
    «/import_topups». Подходит либо выгрузка из CRM «Лог ручных начислений»
    как есть (бот сам посчитает средние), либо файл из двух столбцов:
    телефон и сумма в месяц. Работает только для администраторов и только
    с этой подписью — случайно присланный файл ничего не сломает."""
    if not is_admin(message.from_user.id):
        return

    caption = (message.caption or "").strip().lower()
    if "import_topups" not in caption:
        return

    filename = (message.document.file_name or "").lower()
    if not filename.endswith((".xlsx", ".xlsm")):
        await message.answer("Нужен файл Excel (.xlsx) — пришли выгрузку в этом формате с той же подписью.")
        return

    try:
        from openpyxl import load_workbook
    except ImportError:
        await message.answer(
            "На сервере не установлен openpyxl — добавь его в requirements.txt и передеплой, "
            "потом пришли файл ещё раз."
        )
        return

    buf = BytesIO()
    await message.bot.download(message.document, destination=buf)
    buf.seek(0)

    try:
        wb = load_workbook(buf, read_only=True, data_only=True)
        ws = wb.active
    except Exception as e:  # noqa: BLE001 — сообщаем причину и не роняем бота
        await message.answer(f"Не получилось открыть файл: {e}")
        return

    kind, stats, info = parse_topups(ws.iter_rows(values_only=True))
    wb.close()
    if not stats:
        await message.answer(
            "В файле не нашлось ни одного телефона с суммой. Нужна либо выгрузка из CRM "
            "(с колонками «Телефон», «Баланс», «Дата»), либо два столбца: телефон и сумма в месяц."
        )
        return

    db = SessionLocal()
    try:
        # Полный лог из CRM заменяет статистику целиком; файл «телефон + сумма» — только дополняет.
        in_app, reset = save_phone_stats(db, stats, replace_all=kind in ("operations_log", "crm_log"))
    finally:
        db.close()

    counts = {name: 0 for name in TIER_LABELS}
    for st in stats.values():
        counts[get_tier(st["avg"])] += 1

    fmt = lambda n: f"{n:,}".replace(",", " ")  # noqa: E731
    lines = ["✅ Пополнения загружены."]
    if kind == "operations_log":
        lines.append(
            f"Формат: лог финансовых операций за {info['period_start']:%d.%m.%Y} — {info['period_end']:%d.%m.%Y}. "
            f"Пополнений учтено: {fmt(info['used'])}."
        )
    elif kind == "crm_log":
        lines.append(
            f"Формат: выгрузка CRM за {info['period_start']:%d.%m.%Y} — {info['period_end']:%d.%m.%Y}. "
            f"Пополнений учтено: {fmt(info['used'])}, строк с бонусами пропущено: {fmt(info['bonus_rows'])}."
        )
    else:
        lines.append("Формат: телефон + сумма в месяц.")
    lines += [
        f"Гостей в файле: {fmt(len(stats))}. Уже в приложении — уровень обновлён сразу: {fmt(in_app)}. "
        f"Остальные увидят уровень, как только привяжут номер.",
        "",
        f"Premium: {counts['premium']} · Gold: {counts['gold']} · Silver: {counts['silver']} · "
        f"без статуса (ниже {fmt(SILVER_QUALIFY_FROM)} ₽): {fmt(counts['none'])}",
    ]
    if reset:
        lines.append(f"Нет в новой выгрузке — статус обнулён: {fmt(reset)} (перестали пополнять).")
    if info.get("bad"):
        lines.append(f"Не удалось прочитать строк: {info['bad']} (обычно это заголовок).")
    await message.answer("\n".join(lines))


@router.message(Command("export"))
async def cmd_export(message: Message):
    """Присылает Excel со ВСЕМИ пользователями приложения — без лимита в 20."""
    if not is_admin(message.from_user.id):
        return

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from aiogram.types import BufferedInputFile
    from datetime import timedelta
    from .models import TIER_CASHBACK

    msk = lambda dt: (dt + timedelta(hours=3)).replace(microsecond=0) if dt else None  # noqa: E731

    db = SessionLocal()
    try:
        clients = db.query(Client).order_by(Client.monthly_topup.desc()).all()
        topups = load_topup_map(db)
    finally:
        db.close()

    if not clients:
        await message.answer("Пользователей пока нет.")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "Пользователи"
    headers = ["Имя в Telegram", "Telegram ID", "Телефон", "Номер подтверждён", "Уровень", "Кэшбэк, %",
               "Среднее в месяц, ₽", "Баланс бонусов", "Согласие на ПДн (МСК)",
               "Первый вход (МСК)", "Последний визит (МСК)", "Найден в CRM", "Номер в CRM (если другой)",
               "Отписан от рассылки"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F1F1F")
    for c in clients:
        key = topup_key(c)
        topup = topups.get(key, 0) if key else 0
        tier = get_tier(topup)
        ws.append([
            c.tg_name, c.tg_id, c.phone, "да" if c.phone_verified else "нет", TIER_LABELS[tier], TIER_CASHBACK[tier],
            round(topup), c.balance or 0,
            msk(c.consent_at), msk(c.created_at), msk(c.last_seen_at),
            ("да" if key in topups else "нет") if key else "",
            f"+7{c.crm_phone_normalized}" if c.crm_phone_normalized else "",
            "да" if c.broadcast_opt_out else "",
        ])
    for col, width in zip("ABCDEFGHIJKLMN", (22, 14, 18, 12, 13, 10, 16, 14, 20, 20, 20, 12, 18, 12)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2, min_col=9, max_col=11):
        for cell in row:
            cell.number_format = "DD.MM.YYYY HH:MM"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    buf = BytesIO()
    wb.save(buf)
    with_phone = sum(1 for c in clients if c.phone)
    stamp = (datetime.utcnow() + timedelta(hours=3)).strftime("%Y-%m-%d")
    await message.answer_document(
        BufferedInputFile(buf.getvalue(), filename=f"colizeum-users-{stamp}.xlsx"),
        caption=f"Пользователей: {len(clients)}, из них привязали телефон: {with_phone}.",
    )


@router.message(Command("link_crm"))
async def cmd_link_crm(message: Message, command: CommandObject):
    """Формат: /link_crm +7(номер в Telegram) +7(номер в CRM)
    Когда гость в Telegram сидит на одном номере, а в CRM клуба записан на
    другой, — связываем: уровень будет считаться по номеру из CRM.
    /link_crm +7(номер в Telegram) — без второго номера убирает связь."""
    if not is_admin(message.from_user.id):
        return
    parts = (command.args or "").split()
    if len(parts) not in (1, 2):
        await message.answer("Формат: /link_crm +7(номер в Telegram) +7(номер в CRM)\n"
                             "Убрать связь: /link_crm +7(номер в Telegram)")
        return
    tg_norm = normalize_phone(parts[0])
    crm_norm = normalize_phone(parts[1]) if len(parts) == 2 else None
    if len(tg_norm) < 10 or (crm_norm is not None and len(crm_norm) < 10):
        await message.answer("Не похоже на номер телефона — нужно 10 цифр после +7")
        return

    db = SessionLocal()
    try:
        client = db.query(Client).filter(Client.phone_normalized == tg_norm,
                                         Client.phone_verified == True).first()  # noqa: E712
        if client is None:
            await message.answer(f"Гость с подтверждённым номером +7{tg_norm} не найден. "
                                 "Сначала он должен нажать «Поделиться номером» в кабинете.")
            return
        client.crm_phone_normalized = crm_norm if crm_norm != tg_norm else None
        db.commit()
        topup, found = lookup_topup(db, client)
        client.monthly_topup = topup
        db.commit()
        name = client.tg_name or "гость"
    finally:
        db.close()

    if crm_norm is None:
        await message.answer(f"Связь убрана. {name}: уровень снова считается по номеру +7{tg_norm} — "
                             f"{TIER_LABELS[get_tier(topup)]}.")
        return
    status = (f"в CRM {round(topup):,} ₽/мес → {TIER_LABELS[get_tier(topup)]}".replace(",", " ")
              if found else "но этого номера тоже нет в статистике CRM — проверь номер (/check)")
    await message.answer(f"Готово: {name} (+7{tg_norm}) связан с номером CRM +7{crm_norm} — {status}.")


# ---------- рассылка ----------

BROADCAST_SEGMENTS = {
    "all": "всем",
    "status": "всем со статусом (Silver, Gold, Premium)",
    "premium": "Premium",
    "gold": "Gold",
    "silver": "Silver",
    "none": "без статуса",
    "nophone": "не подтвердившим номер",
}
_pending_broadcasts: dict[str, dict] = {}   # ключ — токен конкретного превью
_bg_tasks: set = set()                       # держим ссылки, чтобы рассылку не прервал сборщик мусора


def broadcast_recipients(db, segment: str) -> list[int]:
    """Кому уйдёт рассылка: только гости, давшие согласие и не отписавшиеся."""
    clients = db.query(Client).filter(
        Client.consent_at.isnot(None),
        or_(Client.broadcast_opt_out == False, Client.broadcast_opt_out.is_(None)),  # noqa: E712
    ).all()
    if segment == "all":
        return [c.tg_id for c in clients]
    if segment == "nophone":
        return [c.tg_id for c in clients if not c.phone_verified]
    topups = load_topup_map(db)
    result = []
    for c in clients:
        key = topup_key(c)
        tier = get_tier(topups.get(key, 0) if key else 0)
        if (segment == "status" and tier != "none") or tier == segment:
            result.append(c.tg_id)
    return result


BROADCAST_HELP = (
    "Как сделать рассылку:\n"
    "1. Напиши боту сообщение, которое хочешь отправить, — можно с фото, жирным шрифтом, ссылками.\n"
    "2. Ответь на него (свайп → «Ответить») командой /broadcast — уйдёт всем.\n"
    "   Только части гостей: /broadcast gold, /broadcast status и т.д.\n"
    "3. Бот покажет, сколько получателей, и попросит подтвердить.\n\n"
    "Короткий текст можно сразу: /broadcast Текст сообщения\n\n"
    "Сегменты: " + ", ".join(f"{k} — {v}" for k, v in BROADCAST_SEGMENTS.items()) + "\n\n"
    "Получают только гости, которые приняли правила в кабинете и не отписались (/stop)."
)


@router.message(Command("broadcast"))
async def cmd_broadcast(message: Message, command: CommandObject):
    if not is_admin(message.from_user.id):
        return
    args = (command.args or "").strip()
    reply = message.reply_to_message

    if reply:
        segment = args.lower() or "all"
        if segment not in BROADCAST_SEGMENTS:
            await message.answer("Не знаю такой сегмент.\n\n" + BROADCAST_HELP)
            return
        source = {"kind": "copy", "chat_id": message.chat.id, "message_id": reply.message_id}
    else:
        if not args:
            await message.answer(BROADCAST_HELP)
            return
        first, _, rest = args.partition(" ")
        if first.lower() in BROADCAST_SEGMENTS and rest.strip():
            segment, text = first.lower(), rest.strip()
        else:
            segment, text = "all", args
        source = {"kind": "text", "text": text}

    db = SessionLocal()
    try:
        count = len(broadcast_recipients(db, segment))
    finally:
        db.close()
    if count == 0:
        await message.answer(f"Получателей нет (сегмент: {BROADCAST_SEGMENTS[segment]}).")
        return

    token = secrets.token_hex(4)
    _pending_broadcasts[token] = {"segment": segment, "source": source, "admin": message.from_user.id}
    preview = "" if source["kind"] == "copy" else f"\n\nТекст:\n{source['text']}"
    await message.answer(
        f"Рассылка: {BROADCAST_SEGMENTS[segment]}. Получателей: {count}.{preview}\n\nОтправляем?",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[[
            InlineKeyboardButton(text=f"✅ Отправить ({count})", callback_data=f"bc:go:{token}"),
            InlineKeyboardButton(text="Отмена", callback_data=f"bc:cancel:{token}"),
        ]]),
    )


@router.callback_query(F.data.startswith("bc:"))
async def cb_broadcast(callback: CallbackQuery):
    if not is_admin(callback.from_user.id):
        await callback.answer()
        return
    _, action, token = (callback.data.split(":") + ["", ""])[:3]
    pending = _pending_broadcasts.pop(token, None)
    if action == "cancel" or pending is None:
        await callback.answer("Отменено" if pending else "Нечего отправлять")
        await callback.message.edit_text("Рассылка отменена." if pending else "Эта рассылка уже неактуальна.")
        return
    await callback.answer("Отправляю")
    await callback.message.edit_text("⏳ Рассылка запущена — пришлю отчёт, когда закончу.")
    task = asyncio.create_task(run_broadcast(callback.bot, pending, callback.from_user.id))
    _bg_tasks.add(task)
    task.add_done_callback(_bg_tasks.discard)


async def _send_one(bot: Bot, tg_id: int, source: dict):
    if source["kind"] == "copy":
        await bot.copy_message(tg_id, from_chat_id=source["chat_id"], message_id=source["message_id"])
    else:
        await bot.send_message(tg_id, source["text"], disable_web_page_preview=False)


async def run_broadcast(bot: Bot, pending: dict, report_chat_id: int):
    db = SessionLocal()
    try:
        recipients = broadcast_recipients(db, pending["segment"])
    finally:
        db.close()

    sent = blocked = failed = 0
    for tg_id in recipients:
        for attempt in range(2):
            try:
                await _send_one(bot, tg_id, pending["source"])
                sent += 1
                break
            except TelegramRetryAfter as e:      # Telegram просит притормозить
                await asyncio.sleep(e.retry_after + 1)
                if attempt == 1:
                    failed += 1
            except TelegramForbiddenError:       # гость заблокировал бота
                blocked += 1
                break
            except TelegramBadRequest:           # чат не найден — гость ни разу не запускал бота
                failed += 1
                break
            except Exception as e:  # noqa: BLE001
                logger.warning("Рассылка: не отправлено %s: %s", tg_id, e)
                failed += 1
                break
        await asyncio.sleep(0.05)   # ~20 сообщений в секунду — в пределах лимитов Telegram

    try:
        await bot.send_message(
            report_chat_id,
            f"✅ Рассылка завершена ({BROADCAST_SEGMENTS[pending['segment']]}).\n"
            f"Доставлено: {sent}\n"
            f"Заблокировали бота: {blocked}\n"
            f"Не доставлено по другим причинам: {failed} (обычно гость открывал кабинет, "
            "но ни разу не нажимал «Старт» в чате с ботом)")
    except Exception as e:  # noqa: BLE001
        logger.warning("Не удалось отправить отчёт о рассылке: %s", e)


GAME_NAMES = {"spin": "Рулетка", "dice": "Кости", "probability": "Дайс",
              "crash": "Crash", "slots": "Слот", "hilo": "Больше-меньше"}


@router.message(Command("games"))
async def cmd_games(message: Message):
    """Экономика игр: сколько бонусов роздано бесплатно и сколько клуб
    «забрал» в играх на ставки — за 7 дней и за всё время."""
    if not is_admin(message.from_user.id):
        return
    fmt = lambda n: f"{int(n):,}".replace(",", " ")  # noqa: E731
    db = SessionLocal()
    try:
        def collect(since=None):
            q = db.query(GamePlay.game, GamePlay.mode, func.count(GamePlay.id),
                         func.coalesce(func.sum(GamePlay.stake), 0), func.coalesce(func.sum(GamePlay.payout), 0))
            if since is not None:
                q = q.filter(GamePlay.created_at >= since)
            return q.group_by(GamePlay.game, GamePlay.mode).all()

        lines = []
        for title, since in (("За 7 дней", datetime.utcnow() - timedelta(days=7)), ("За всё время", None)):
            rows = collect(since)
            lines.append(f"📊 {title}")
            free_total = bet_net = 0
            for game, mode, n, staked, paid in sorted(rows, key=lambda r: (r[1], r[0])):
                name = GAME_NAMES.get(game, game)
                if mode == "free":
                    free_total += paid
                    lines.append(f"  {name} (бесплатно): {fmt(n)} игр, роздано {fmt(paid)}")
                else:
                    net = staked - paid
                    bet_net += net
                    lines.append(f"  {name} (на бонусы): {fmt(n)} ставок, поставлено {fmt(staked)}, "
                                 f"выплачено {fmt(paid)}, клубу {'+' if net >= 0 else ''}{fmt(net)}")
            if not rows:
                lines.append("  игр пока не было")
            lines.append(f"  Итого: роздано бесплатно {fmt(free_total)}, "
                         f"сгорело в ставках {'+' if bet_net >= 0 else ''}{fmt(bet_net)}")
            lines.append("")
        await message.answer("\n".join(lines).strip())
    finally:
        db.close()


@router.message(Command("admin"))
async def cmd_admin(message: Message):
    """Шпаргалка по всем админ-командам."""
    if not is_admin(message.from_user.id):
        return
    await message.answer(
        "Команды администратора:\n\n"
        "Статистика\n"
        "/stats — сводка: гости, уровни, найдены ли в CRM\n"
        "/clients — последние 20 гостей\n"
        "/export — Excel со всеми гостями\n"
        "/games — экономика игр: сколько роздано и сгорело\n"
        "/check +7… — всё по номеру (CRM, аккаунты, уровень)\n\n"
        "Уровни\n"
        "Файл Excel с подписью /import_topups — загрузить пополнения из CRM\n"
        "/set_topup +7… 12000 — вручную задать среднее в месяц\n"
        "/link_crm +7(Telegram) +7(CRM) — если в CRM гость на другом номере\n\n"
        "Бонусы\n"
        "/add_bonus +7… 150 — начислить в приложении\n"
        "/spend +7… 150 — списать в приложении\n\n"
        "Рассылка\n"
        "/broadcast — как сделать рассылку (ответом на сообщение)\n\n"
        "Акции\n"
        "/promotions — список, /add_promo Название | Описание | ссылка"
    )


def build_bot_and_dispatcher() -> tuple[Bot, Dispatcher]:
    bot_token = os.getenv("BOT_TOKEN", "")
    if not bot_token:
        raise RuntimeError("Не задан BOT_TOKEN — возьми токен у @BotFather и добавь в .env")

    bot = Bot(token=bot_token)
    dp = Dispatcher()
    dp.include_router(router)
    return bot, dp
