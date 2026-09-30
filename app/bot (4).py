"""
Telegram-бот клуба COLIZEUM.

Что делает:
- /start — присылает кнопку, открывающую мини-приложение (личный кабинет).
- Админ-команды (доступны только tg_id из ADMIN_IDS) — временный способ
  вести данные о пополнениях и акциях, пока нет автоматической синхронизации
  с CRM/POS клуба. Как только появится выгрузка или API, эти команды можно
  будет заменить на автоматический импорт, ничего в остальном коде менять
  не придётся.

Бот работает поверх той же базы данных, что и веб-API мини-приложения
(см. main.py) — это один процесс, поэтому отдельный сервис под бота
разворачивать не нужно.
"""
import os
from datetime import datetime
from io import BytesIO

from aiogram import Bot, Dispatcher, F, Router
from aiogram.filters import Command, CommandObject
from aiogram.types import Message, WebAppInfo, InlineKeyboardMarkup, InlineKeyboardButton

from .database import SessionLocal
from .models import (
    Client, Promotion, PhoneStat, get_tier, normalize_phone,
    TIER_LABELS, SILVER_QUALIFY_FROM,
)
from .topups import parse_topups

WEBAPP_URL = os.getenv("WEBAPP_URL", "")
ADMIN_IDS = {
    int(x) for x in os.getenv("ADMIN_IDS", "").replace(" ", "").split(",") if x
}

router = Router()


def is_admin(tg_id: int) -> bool:
    return tg_id in ADMIN_IDS


def humanize_ago(dt) -> str:
    """Переводит момент времени в «X назад» на русском — для списка гостей."""
    if not dt:
        return "не заходил(а)"
    delta = datetime.utcnow() - dt
    minutes = int(delta.total_seconds() // 60)
    if minutes < 1:
        return "только что"
    if minutes < 60:
        return f"{minutes} мин назад"
    hours = minutes // 60
    if hours < 24:
        return f"{hours} ч назад"
    days = hours // 24
    return f"{days} дн назад"


@router.message(Command("start"))
async def cmd_start(message: Message):
    if not WEBAPP_URL:
        await message.answer(
            "Мини-приложение ещё не подключено (не задан WEBAPP_URL). "
            "Как только опубликуешь webapp, добавь адрес в настройки — и эта кнопка заработает."
        )
        return

    keyboard = InlineKeyboardMarkup(
        inline_keyboard=[[
            InlineKeyboardButton(text="Открыть личный кабинет", web_app=WebAppInfo(url=WEBAPP_URL))
        ]]
    )
    await message.answer(
        "Добро пожаловать в COLIZEUM.\n\n"
        "Здесь — твой уровень, акции клуба и мини-игры с призами. "
        "Открывай, когда удобно.",
        reply_markup=keyboard,
    )


def save_phone_stats(db, stats: dict, replace_all: bool = False) -> tuple[int, int]:
    """Сохраняет средние пополнения по телефонам и сразу проставляет их
    гостям, которые уже привязали номер.

    replace_all=True — для полной выгрузки из CRM: статистика заменяется
    целиком. Телефоны, которых нет в новой выгрузке (гость перестал ходить),
    обнуляются, чтобы старый статус не висел вечно.
    Возвращает (скольким гостям в приложении обновлён уровень, скольким обнулён)."""
    now = datetime.utcnow()
    reset = 0
    if replace_all:
        db.query(PhoneStat).delete(synchronize_session=False)
        stale = db.query(Client).filter(Client.phone_normalized.isnot(None),
                                        Client.monthly_topup > 0).all()
        for client in stale:
            if client.phone_normalized not in stats:
                client.monthly_topup = 0.0
                reset += 1
        db.flush()
    phones = list(stats)
    chunks = [phones[i:i + 500] for i in range(0, len(phones), 500)]  # лимит параметров в SQL
    existing = {}
    for chunk in chunks:
        for s in db.query(PhoneStat).filter(PhoneStat.phone_normalized.in_(chunk)).all():
            existing[s.phone_normalized] = s
    for phone, st in stats.items():
        row = existing.get(phone)
        if row is None:
            db.add(PhoneStat(phone_normalized=phone, avg_monthly=st["avg"],
                             total=st.get("total"), updated_at=now))
        else:
            row.avg_monthly = st["avg"]
            row.total = st.get("total")
            row.updated_at = now

    updated = 0
    for chunk in chunks:
        for client in db.query(Client).filter(Client.phone_normalized.in_(chunk)).all():
            client.monthly_topup = stats[client.phone_normalized]["avg"]
            updated += 1
    db.commit()
    return updated, reset


@router.message(Command("set_topup"))
async def cmd_set_topup(message: Message, command: CommandObject):
    """Формат: /set_topup +79991234567 12000
    Вручную задаёт средние пополнения в месяц по номеру. Работает и для
    гостей, которые ещё не заходили в приложение — уровень появится у них
    при привязке номера."""
    if not is_admin(message.from_user.id):
        return

    parts = (command.args or "").split()
    if len(parts) != 2:
        await message.answer("Формат: /set_topup +79991234567 12000")
        return

    phone, amount_raw = parts
    try:
        amount = float(amount_raw.replace(",", "."))
    except ValueError:
        await message.answer("Сумма должна быть числом, например 12000")
        return

    norm = normalize_phone(phone)
    if len(norm) < 10:
        await message.answer("Не похоже на номер телефона — нужно 10 цифр после +7")
        return

    db = SessionLocal()
    try:
        updated, _ = save_phone_stats(db, {norm: {"avg": amount, "total": None}})
    finally:
        db.close()

    tier = TIER_LABELS[get_tier(amount)]
    tail = "Гость уже в приложении — уровень обновлён." if updated else \
        "Гость ещё не привязал номер — уровень появится сразу при привязке."
    await message.answer(f"Готово: {phone} — {amount:,.0f} ₽ в месяц, уровень {tier}.\n{tail}".replace(",", " "))


@router.message(Command("add_promo"))
async def cmd_add_promo(message: Message, command: CommandObject):
    """Формат: /add_promo Название | Описание | ссылка (необязательно)
    Акции видны всем гостям клуба независимо от уровня."""
    if not is_admin(message.from_user.id):
        return

    if not command.args:
        await message.answer("Формат: /add_promo Название | Описание | ссылка (необязательно)")
        return

    parts = [p.strip() for p in command.args.split("|")]
    if len(parts) < 2:
        await message.answer("Формат: /add_promo Название | Описание | ссылка (необязательно)")
        return

    title, description = parts[0], parts[1]
    link = parts[2] if len(parts) > 2 and parts[2] else None

    db = SessionLocal()
    try:
        promo = Promotion(title=title, description=description, link=link)
        db.add(promo)
        db.commit()
        await message.answer(f"Акция добавлена: «{title}»" + (f"\nСсылка: {link}" if link else ""))
    finally:
        db.close()


@router.message(Command("promotions"))
async def cmd_list_promotions(message: Message):
    if not is_admin(message.from_user.id):
        return

    db = SessionLocal()
    try:
        promos = db.query(Promotion).filter(Promotion.active == True).all()  # noqa: E712
        if not promos:
            await message.answer("Активных акций пока нет. Добавь: /add_promo Название | Описание")
            return
        lines = [f"#{p.id} {p.title} — {p.description}" for p in promos]
        await message.answer("\n".join(lines))
    finally:
        db.close()


@router.message(Command("stats"))
async def cmd_stats(message: Message):
    """Сводка: сколько всего гостей в базе, разбивка по уровням,
    сколько заходило за последние сутки и за последнюю неделю."""
    if not is_admin(message.from_user.id):
        return

    db = SessionLocal()
    try:
        clients = db.query(Client).all()
        total = len(clients)
        if total == 0:
            await message.answer("Гостей пока нет — база пустая.")
            return

        by_tier = {name: 0 for name in TIER_LABELS}
        active_24h = 0
        active_7d = 0
        now = datetime.utcnow()
        for c in clients:
            by_tier[get_tier(c.monthly_topup)] += 1
            if c.last_seen_at:
                delta_hours = (now - c.last_seen_at).total_seconds() / 3600
                if delta_hours <= 24:
                    active_24h += 1
                if delta_hours <= 24 * 7:
                    active_7d += 1

        lines = [
            f"Всего гостей в базе: {total}",
            f"Заходили за 24 часа: {active_24h}",
            f"Заходили за 7 дней: {active_7d}",
            "",
            "По уровням:",
            f"  Premium — {by_tier['premium']}",
            f"  Gold — {by_tier['gold']}",
            f"  Silver — {by_tier['silver']}",
            f"  Без статуса — {by_tier['none']}",
        ]
        await message.answer("\n".join(lines))
    finally:
        db.close()


@router.message(Command("clients"))
async def cmd_clients(message: Message):
    """Список последних гостей по времени визита — кто и когда заходил.
    Показывает 20 самых недавних; для точной сверки сумм используй выгрузку из CRM."""
    if not is_admin(message.from_user.id):
        return

    db = SessionLocal()
    try:
        clients = (
            db.query(Client)
            .order_by(Client.last_seen_at.desc().nullslast())
            .limit(20)
            .all()
        )
        if not clients:
            await message.answer("Гостей пока нет — база пустая.")
            return

        lines = ["Последние 20 гостей (по времени визита):", ""]
        for c in clients:
            name = c.tg_name or "без имени"
            phone = c.phone or "телефон не привязан"
            tier = TIER_LABELS[get_tier(c.monthly_topup)]
            lines.append(f"{name} — {phone} — {tier} — {humanize_ago(c.last_seen_at)}")
        await message.answer("\n".join(lines))
    finally:
        db.close()


@router.message(Command("spend"))
async def cmd_spend(message: Message, command: CommandObject):
    """Формат: /spend +79991234567 150
    Списывает бонусы с баланса гостя, когда он расплачивается ими на кассе."""
    if not is_admin(message.from_user.id):
        return

    if not command.args:
        await message.answer("Формат: /spend +79991234567 150")
        return

    parts = command.args.split()
    if len(parts) != 2:
        await message.answer("Формат: /spend +79991234567 150")
        return

    phone, amount_raw = parts
    try:
        amount = int(float(amount_raw))
    except ValueError:
        await message.answer("Сумма должна быть числом, например 150")
        return

    if amount <= 0:
        await message.answer("Сумма списания должна быть больше нуля")
        return

    db = SessionLocal()
    try:
        client = db.query(Client).filter(Client.phone_normalized == normalize_phone(phone)).first()
        if not client:
            await message.answer(f"Клиент с номером {phone} не найден — он ещё не заходил в приложение.")
            return
        current = client.balance or 0
        if current < amount:
            await message.answer(
                f"Недостаточно бонусов: на балансе {current}, а списать нужно {amount}."
            )
            return
        client.balance = current - amount
        db.commit()
        await message.answer(f"Списано {amount} бонусов. Остаток на балансе: {client.balance}.")
    finally:
        db.close()


@router.message(Command("add_bonus"))
async def cmd_add_bonus(message: Message, command: CommandObject):
    """Формат: /add_bonus +79991234567 150
    Начисляет бонусы на баланс в приложении — например, вернуть бонусы
    по отклонённой заявке на списание или начислить вручную."""
    if not is_admin(message.from_user.id):
        return

    parts = (command.args or "").split()
    if len(parts) != 2:
        await message.answer("Формат: /add_bonus +79991234567 150")
        return

    phone, amount_raw = parts
    try:
        amount = int(float(amount_raw))
    except ValueError:
        await message.answer("Сумма должна быть числом, например 150")
        return
    if amount <= 0:
        await message.answer("Сумма должна быть больше нуля")
        return

    db = SessionLocal()
    try:
        client = db.query(Client).filter(Client.phone_normalized == normalize_phone(phone)).first()
        if not client:
            await message.answer(
                f"Клиент с номером {phone} не найден — номер должен быть привязан в приложении."
            )
            return
        client.balance = (client.balance or 0) + amount
        db.commit()
        await message.answer(f"Начислено {amount} бонусов. Баланс в приложении: {client.balance}.")
    finally:
        db.close()


@router.message(F.document)
async def handle_document(message: Message):
    """Загрузка пополнений. Прикрепи в Telegram файл Excel (.xlsx) с подписью
    «/import_topups». Подходит либо выгрузка из CRM «Лог ручных начислений»
    как есть (бот сам посчитает средние), либо файл из двух столбцов:
    телефон и сумма в месяц. Работает только для администраторов и только
    с этой подписью — случайно присланный файл ничего не сломает."""
    if not is_admin(message.from_user.id):
        return

    caption = (message.caption or "").strip().lower()
    if "import_topups" not in caption:
        return

    filename = (message.document.file_name or "").lower()
    if not filename.endswith((".xlsx", ".xlsm")):
        await message.answer("Нужен файл Excel (.xlsx) — пришли выгрузку в этом формате с той же подписью.")
        return

    try:
        from openpyxl import load_workbook
    except ImportError:
        await message.answer(
            "На сервере не установлен openpyxl — добавь его в requirements.txt и передеплой, "
            "потом пришли файл ещё раз."
        )
        return

    buf = BytesIO()
    await message.bot.download(message.document, destination=buf)
    buf.seek(0)

    try:
        wb = load_workbook(buf, read_only=True, data_only=True)
        ws = wb.active
    except Exception as e:  # noqa: BLE001 — сообщаем причину и не роняем бота
        await message.answer(f"Не получилось открыть файл: {e}")
        return

    kind, stats, info = parse_topups(ws.iter_rows(values_only=True))
    wb.close()
    if not stats:
        await message.answer(
            "В файле не нашлось ни одного телефона с суммой. Нужна либо выгрузка из CRM "
            "(с колонками «Телефон», «Баланс», «Дата»), либо два столбца: телефон и сумма в месяц."
        )
        return

    db = SessionLocal()
    try:
        # Полный лог из CRM заменяет статистику целиком; файл «телефон + сумма» — только дополняет.
        in_app, reset = save_phone_stats(db, stats, replace_all=kind in ("operations_log", "crm_log"))
    finally:
        db.close()

    counts = {name: 0 for name in TIER_LABELS}
    for st in stats.values():
        counts[get_tier(st["avg"])] += 1

    fmt = lambda n: f"{n:,}".replace(",", " ")  # noqa: E731
    lines = ["✅ Пополнения загружены."]
    if kind == "operations_log":
        lines.append(
            f"Формат: лог финансовых операций за {info['period_start']:%d.%m.%Y} — {info['period_end']:%d.%m.%Y}. "
            f"Пополнений учтено: {fmt(info['used'])}."
        )
    elif kind == "crm_log":
        lines.append(
            f"Формат: выгрузка CRM за {info['period_start']:%d.%m.%Y} — {info['period_end']:%d.%m.%Y}. "
            f"Пополнений учтено: {fmt(info['used'])}, строк с бонусами пропущено: {fmt(info['bonus_rows'])}."
        )
    else:
        lines.append("Формат: телефон + сумма в месяц.")
    lines += [
        f"Гостей в файле: {fmt(len(stats))}. Уже в приложении — уровень обновлён сразу: {fmt(in_app)}. "
        f"Остальные увидят уровень, как только привяжут номер.",
        "",
        f"Premium: {counts['premium']} · Gold: {counts['gold']} · Silver: {counts['silver']} · "
        f"без статуса (ниже {fmt(SILVER_QUALIFY_FROM)} ₽): {fmt(counts['none'])}",
    ]
    if reset:
        lines.append(f"Нет в новой выгрузке — статус обнулён: {fmt(reset)} (перестали пополнять).")
    if info.get("bad"):
        lines.append(f"Не удалось прочитать строк: {info['bad']} (обычно это заголовок).")
    await message.answer("\n".join(lines))


@router.message(Command("export"))
async def cmd_export(message: Message):
    """Присылает Excel со ВСЕМИ пользователями приложения — без лимита в 20."""
    if not is_admin(message.from_user.id):
        return

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from aiogram.types import BufferedInputFile
    from datetime import timedelta
    from .models import TIER_CASHBACK

    msk = lambda dt: (dt + timedelta(hours=3)).replace(microsecond=0) if dt else None  # noqa: E731

    db = SessionLocal()
    try:
        clients = db.query(Client).order_by(Client.monthly_topup.desc()).all()
    finally:
        db.close()

    if not clients:
        await message.answer("Пользователей пока нет.")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "Пользователи"
    headers = ["Имя в Telegram", "Telegram ID", "Телефон", "Уровень", "Кэшбэк, %",
               "Среднее в месяц, ₽", "Баланс бонусов", "Согласие на ПДн (МСК)",
               "Первый вход (МСК)", "Последний визит (МСК)"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F1F1F")
    for c in clients:
        tier = get_tier(c.monthly_topup or 0)
        ws.append([
            c.tg_name, c.tg_id, c.phone, TIER_LABELS[tier], TIER_CASHBACK[tier],
            round(c.monthly_topup or 0), c.balance or 0,
            msk(c.consent_at), msk(c.created_at), msk(c.last_seen_at),
        ])
    for col, width in zip("ABCDEFGHIJ", (22, 14, 18, 13, 10, 16, 14, 20, 20, 20)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2, min_col=8, max_col=10):
        for cell in row:
            cell.number_format = "DD.MM.YYYY HH:MM"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    buf = BytesIO()
    wb.save(buf)
    with_phone = sum(1 for c in clients if c.phone)
    stamp = (datetime.utcnow() + timedelta(hours=3)).strftime("%Y-%m-%d")
    await message.answer_document(
        BufferedInputFile(buf.getvalue(), filename=f"colizeum-users-{stamp}.xlsx"),
        caption=f"Пользователей: {len(clients)}, из них привязали телефон: {with_phone}.",
    )


def build_bot_and_dispatcher() -> tuple[Bot, Dispatcher]:
    bot_token = os.getenv("BOT_TOKEN", "")
    if not bot_token:
        raise RuntimeError("Не задан BOT_TOKEN — возьми токен у @BotFather и добавь в .env")

    bot = Bot(token=bot_token)
    dp = Dispatcher()
    dp.include_router(router)
    return bot, dp
