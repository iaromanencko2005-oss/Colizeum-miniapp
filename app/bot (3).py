"""
Telegram-бот клуба COLIZEUM.

Что делает:
- /start — присылает кнопку, открывающую мини-приложение (личный кабинет).
- Админ-команды (доступны только tg_id из ADMIN_IDS) — временный способ
  вести данные о пополнениях и акциях, пока нет автоматической синхронизации
  с CRM/POS клуба. Как только появится выгрузка или API, эти команды можно
  будет заменить на автоматический импорт, ничего в остальном коде менять
  не придётся.

Бот работает поверх той же базы данных, что и веб-API мини-приложения
(см. main.py) — это один процесс, поэтому отдельный сервис под бота
разворачивать не нужно.
"""
import os
from datetime import datetime
from io import BytesIO

from aiogram import Bot, Dispatcher, F, Router
from aiogram.filters import Command, CommandObject
from aiogram.types import Message, WebAppInfo, InlineKeyboardMarkup, InlineKeyboardButton

from .database import SessionLocal
from .models import (
    Client, Promotion, PhoneStat, get_tier, normalize_phone,
    TIER_LABELS, SILVER_QUALIFY_FROM,
)
from .topups import parse_topups

WEBAPP_URL = os.getenv("WEBAPP_URL", "")
ADMIN_IDS = {
    int(x) for x in os.getenv("ADMIN_IDS", "").replace(" ", "").split(",") if x
}

router = Router()


def is_admin(tg_id: int) -> bool:
    return tg_id in ADMIN_IDS


def humanize_ago(dt) -> str:
    """Переводит момент времени в «X назад» на русском — для списка гостей."""
    if not dt:
        return "не заходил(а)"
    delta = datetime.utcnow() - dt
    minutes = int(delta.total_seconds() // 60)
    if minutes < 1:
        return "только что"
    if minutes < 60:
        return f"{minutes} мин назад"
    hours = minutes // 60
    if hours < 24:
        return f"{hours} ч назад"
    days = hours // 24
    return f"{days} дн назад"


@router.message(Command("start"))
async def cmd_start(message: Message):
    if not WEBAPP_URL:
        await message.answer(
            "Мини-приложение ещё не подключено (не задан WEBAPP_URL). "
            "Как только опубликуешь webapp, добавь адрес в настройки — и эта кнопка заработает."
        )
        return

    keyboard = InlineKeyboardMarkup(
        inline_keyboard=[[
            InlineKeyboardButton(text="Открыть личный кабинет", web_app=WebAppInfo(url=WEBAPP_URL))
        ]]
    )
    await message.answer(
        "Добро пожаловать в COLIZEUM.\n\n"
        "Здесь — твой уровень, акции клуба и мини-игры с призами. "
        "Открывай, когда удобно.",
        reply_markup=keyboard,
    )


def save_phone_stats(db, stats: dict) -> int:
    """Сохраняет средние пополнения по телефонам (новые добавляет, старые
    обновляет) и сразу проставляет их гостям, которые уже привязали номер.
    Возвращает, скольким гостям в приложении обновлён уровень."""
    now = datetime.utcnow()
    phones = list(stats)
    chunks = [phones[i:i + 500] for i in range(0, len(phones), 500)]  # лимит параметров в SQL
    existing = {}
    for chunk in chunks:
        for s in db.query(PhoneStat).filter(PhoneStat.phone_normalized.in_(chunk)).all():
            existing[s.phone_normalized] = s
    for phone, st in stats.items():
        row = existing.get(phone)
        if row is None:
            db.add(PhoneStat(phone_normalized=phone, avg_monthly=st["avg"],
                             total=st.get("total"), updated_at=now))
        else:
            row.avg_monthly = st["avg"]
            row.total = st.get("total")
            row.updated_at = now

    updated = 0
    for chunk in chunks:
        for client in db.query(Client).filter(Client.phone_normalized.in_(chunk)).all():
            client.monthly_topup = stats[client.phone_normalized]["avg"]
            updated += 1
    db.commit()
    return updated


@router.message(Command("set_topup"))
async def cmd_set_topup(message: Message, command: CommandObject):
    """Формат: /set_topup +79991234567 12000
    Вручную задаёт средние пополнения в месяц по номеру. Работает и для
    гостей, которые ещё не заходили в приложение — уровень появится у них
    при привязке номера."""
    if not is_admin(message.from_user.id):
        return

    parts = (command.args or "").split()
    if len(parts) != 2:
        await message.answer("Формат: /set_topup +79991234567 12000")
        return

    phone, amount_raw = parts
    try:
        amount = float(amount_raw.replace(",", "."))
    except ValueError:
        await message.answer("Сумма должна быть числом, например 12000")
        return

    norm = normalize_phone(phone)
    if len(norm) < 10:
        await message.answer("Не похоже на номер телефона — нужно 10 цифр после +7")
        return

    db = SessionLocal()
    try:
        updated = save_phone_stats(db, {norm: {"avg": amount, "total": None}})
    finally:
        db.close()

    tier = TIER_LABELS[get_tier(amount)]
    tail = "Гость уже в приложении — уровень обновлён." if updated else \
        "Гость ещё не привязал номер — уровень появится сразу при привязке."
    await message.answer(f"Готово: {phone} — {amount:,.0f} ₽ в месяц, уровень {tier}.\n{tail}".replace(",", " "))


@router.message(Command("add_promo"))
async def cmd_add_promo(message: Message, command: CommandObject):
    """Формат: /add_promo Название | Описание | ссылка (необязательно)
    Акции видны всем гостям клуба независимо от уровня."""
    if not is_admin(message.from_user.id):
        return

    if not command.args:
        await message.answer("Формат: /add_promo Название | Описание | ссылка (необязательно)")
        return

    parts = [p.strip() for p in command.args.split("|")]
    if len(parts) < 2:
        await message.answer("Формат: /add_promo Название | Описание | ссылка (необязательно)")
        return

    title, description = parts[0], parts[1]
    link = parts[2] if len(parts) > 2 and parts[2] else None

    db = SessionLocal()
    try:
        promo = Promotion(title=title, description=description, link=link)
        db.add(promo)
        db.commit()
        await message.answer(f"Акция добавлена: «{title}»" + (f"\nСсылка: {link}" if link else ""))
    finally:
        db.close()


@router.message(Command("promotions"))
async def cmd_list_promotions(message: Message):
    if not is_admin(message.from_user.id):
        return

    db = SessionLocal()
    try:
        promos = db.query(Promotion).filter(Promotion.active == True).all()  # noqa: E712
        if not promos:
            await message.answer("Активных акций пока нет. Добавь: /add_promo Название | Описание")
            return
        lines = [f"#{p.id} {p.title} — {p.description}" for p in promos]
        await message.answer("\n".join(lines))
    finally:
        db.close()


@router.message(Command("stats"))
async def cmd_stats(message: Message):
    """Сводка: сколько всего гостей в базе, разбивка по уровням,
    сколько заходило за последние сутки и за последнюю неделю."""
    if not is_admin(message.from_user.id):
        return

    db = SessionLocal()
    try:
        clients = db.query(Client).all()
        total = len(clients)
        if total == 0:
            await message.answer("Гостей пока нет — база пустая.")
            return

        by_tier = {name: 0 for name in TIER_LABELS}
        active_24h = 0
        active_7d = 0
        now = datetime.utcnow()
        for c in clients:
            by_tier[get_tier(c.monthly_topup)] += 1
            if c.last_seen_at:
                delta_hours = (now - c.last_seen_at).total_seconds() / 3600
                if delta_hours <= 24:
                    active_24h += 1
                if delta_hours <= 24 * 7:
                    active_7d += 1

        lines = [
            f"Всего гостей в базе: {total}",
            f"Заходили за 24 часа: {active_24h}",
            f"Заходили за 7 дней: {active_7d}",
            "",
            "По уровням:",
            f"  Premium — {by_tier['premium']}",
            f"  Gold — {by_tier['gold']}",
            f"  Silver — {by_tier['silver']}",
            f"  Без статуса — {by_tier['none']}",
        ]
        await message.answer("\n".join(lines))
    finally:
        db.close()


@router.message(Command("clients"))
async def cmd_clients(message: Message):
    """Список последних гостей по времени визита — кто и когда заходил.
    Показывает 20 самых недавних; для точной сверки сумм используй выгрузку из CRM."""
    if not is_admin(message.from_user.id):
        return

    db = SessionLocal()
    try:
        clients = (
            db.query(Client)
            .order_by(Client.last_seen_at.desc().nullslast())
            .limit(20)
            .all()
        )
        if not clients:
            await message.answer("Гостей пока нет — база пустая.")
            return

        lines = ["Последние 20 гостей (по времени визита):", ""]
        for c in clients:
            name = c.tg_name or "без имени"
            phone = c.phone or "телефон не привязан"
            tier = TIER_LABELS[get_tier(c.monthly_topup)]
            lines.append(f"{name} — {phone} — {tier} — {humanize_ago(c.last_seen_at)}")
        await message.answer("\n".join(lines))
    finally:
        db.close()


@router.message(Command("spend"))
async def cmd_spend(message: Message, command: CommandObject):
    """Формат: /spend +79991234567 150
    Списывает бонусы с баланса гостя, когда он расплачивается ими на кассе."""
    if not is_admin(message.from_user.id):
        return

    if not command.args:
        await message.answer("Формат: /spend +79991234567 150")
        return

    parts = command.args.split()
    if len(parts) != 2:
        await message.answer("Формат: /spend +79991234567 150")
        return

    phone, amount_raw = parts
    try:
        amount = int(float(amount_raw))
    except ValueError:
        await message.answer("Сумма должна быть числом, например 150")
        return

    if amount <= 0:
        await message.answer("Сумма списания должна быть больше нуля")
        return

    db = SessionLocal()
    try:
        client = db.query(Client).filter(Client.phone_normalized == normalize_phone(phone)).first()
        if not client:
            await message.answer(f"Клиент с номером {phone} не найден — он ещё не заходил в приложение.")
            return
        current = client.balance or 0
        if current < amount:
            await message.answer(
                f"Недостаточно бонусов: на балансе {current}, а списать нужно {amount}."
            )
            return
        client.balance = current - amount
        db.commit()
        await message.answer(f"Списано {amount} бонусов. Остаток на балансе: {client.balance}.")
    finally:
        db.close()


@router.message(Command("add_bonus"))
async def cmd_add_bonus(message: Message, command: CommandObject):
    """Формат: /add_bonus +79991234567 150
    Начисляет бонусы на баланс в приложении — например, вернуть бонусы
    по отклонённой заявке на списание или начислить вручную."""
    if not is_admin(message.from_user.id):
        return

    parts = (command.args or "").split()
    if len(parts) != 2:
        await message.answer("Формат: /add_bonus +79991234567 150")
        return

    phone, amount_raw = parts
    try:
        amount = int(float(amount_raw))
    except ValueError:
        await message.answer("Сумма должна быть числом, например 150")
        return
    if amount <= 0:
        await message.answer("Сумма должна быть больше нуля")
        return

    db = SessionLocal()
    try:
        client = db.query(Client).filter(Client.phone_normalized == normalize_phone(phone)).first()
        if not client:
            await message.answer(
                f"Клиент с номером {phone} не найден — номер должен быть привязан в приложении."
            )
            return
        client.balance = (client.balance or 0) + amount
        db.commit()
        await message.answer(f"Начислено {amount} бонусов. Баланс в приложении: {client.balance}.")
    finally:
        db.close()


@router.message(F.document)
async def handle_document(message: Message):
    """Загрузка пополнений. Прикрепи в Telegram файл Excel (.xlsx) с подписью
    «/import_topups». Подходит либо выгрузка из CRM «Лог ручных начислений»
    как есть (бот сам посчитает средние), либо файл из двух столбцов:
    телефон и сумма в месяц. Работает только для администраторов и только
    с этой подписью — случайно присланный файл ничего не сломает."""
    if not is_admin(message.from_user.id):
        return

    caption = (message.caption or "").strip().lower()
    if "import_topups" not in caption:
        return

    filename = (message.document.file_name or "").lower()
    if not filename.endswith((".xlsx", ".xlsm")):
        await message.answer("Нужен файл Excel (.xlsx) — пришли выгрузку в этом формате с той же подписью.")
        return

    try:
        from openpyxl import load_workbook
    except ImportError:
        await message.answer(
            "На сервере не установлен openpyxl — добавь его в requirements.txt и передеплой, "
            "потом пришли файл ещё раз."
        )
        return

    buf = BytesIO()
    await message.bot.download(message.document, destination=buf)
    buf.seek(0)

    try:
        wb = load_workbook(buf, read_only=True, data_only=True)
        ws = wb.active
    except Exception as e:  # noqa: BLE001 — сообщаем причину и не роняем бота
        await message.answer(f"Не получилось открыть файл: {e}")
        return

    kind, stats, info = parse_topups(ws.iter_rows(values_only=True))
    wb.close()
    if not stats:
        await message.answer(
            "В файле не нашлось ни одного телефона с суммой. Нужна либо выгрузка из CRM "
            "(с колонками «Телефон», «Баланс», «Дата»), либо два столбца: телефон и сумма в месяц."
        )
        return

    db = SessionLocal()
    try:
        in_app = save_phone_stats(db, stats)
    finally:
        db.close()

    counts = {name: 0 for name in TIER_LABELS}
    for st in stats.values():
        counts[get_tier(st["avg"])] += 1

    fmt = lambda n: f"{n:,}".replace(",", " ")  # noqa: E731
    lines = ["✅ Пополнения загружены."]
    if kind == "operations_log":
        lines.append(
            f"Формат: лог финансовых операций за {info['period_start']:%d.%m.%Y} — {info['period_end']:%d.%m.%Y}. "
            f"Пополнений учтено: {fmt(info['used'])}."
        )
    elif kind == "crm_log":
        lines.append(
            f"Формат: выгрузка CRM за {info['period_start']:%d.%m.%Y} — {info['period_end']:%d.%m.%Y}. "
            f"Пополнений учтено: {fmt(info['used'])}, строк с бонусами пропущено: {fmt(info['bonus_rows'])}."
        )
    else:
        lines.append("Формат: телефон + сумма в месяц.")
    lines += [
        f"Гостей в файле: {fmt(len(stats))}. Уже в приложении — уровень обновлён сразу: {fmt(in_app)}. "
        f"Остальные увидят уровень, как только привяжут номер.",
        "",
        f"Premium: {counts['premium']} · Gold: {counts['gold']} · Silver: {counts['silver']} · "
        f"без статуса (ниже {fmt(SILVER_QUALIFY_FROM)} ₽): {fmt(counts['none'])}",
    ]
    if info.get("bad"):
        lines.append(f"Не удалось прочитать строк: {info['bad']} (обычно это заголовок).")
    await message.answer("\n".join(lines))


def build_bot_and_dispatcher() -> tuple[Bot, Dispatcher]:
    bot_token = os.getenv("BOT_TOKEN", "")
    if not bot_token:
        raise RuntimeError("Не задан BOT_TOKEN — возьми токен у @BotFather и добавь в .env")

    bot = Bot(token=bot_token)
    dp = Dispatcher()
    dp.include_router(router)
    return bot, dp
